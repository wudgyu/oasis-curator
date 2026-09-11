"""
文档解析器

支持 PDF（pdfplumber）/ Word（python-docx）/ TXT / Markdown，统一输出纯文本。

- PDF 逐页提取并保留页码信息，供切分时标注 chunk 的来源页
- PDF / Word 中的表格转为 Markdown 文本混入正文，使表格内容同样可被检索与引用
  （复制为 Markdown 而不是丢弃，是因为表格答案通常跨行跨列，
   转为文本后 Embedding 才能把"某行数据"与"表头语义"关联起来）
- TXT/Markdown 按文本读取，UTF-8 优先、GBK 回退（中文文档常见编码）
"""

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterator, List, Optional, Tuple

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}

# 表格在正文中的分隔标记，便于人工核对与引用定位
TABLE_MARKER = "[表格 {n}]"


@dataclass
class ParsedDocument:
    """解析后的文档：纯文本 + 元信息"""

    text: str  # 全文纯文本（页间以空行分隔）
    file_name: str  # 原始文件名
    file_type: str  # pdf / word / txt / markdown
    char_count: int = 0  # 总字符数
    page_count: int = 0  # 页数（仅 PDF 有意义）
    pages: List[str] = field(default_factory=list)  # 每页文本（仅 PDF）
    table_count: int = 0  # 提取到的表格数（PDF / Word）


# ---------------------------------------------------------------------------
# 表格 → Markdown
# ---------------------------------------------------------------------------


def table_to_markdown(rows: List[List[Optional[str]]]) -> str:
    """
    二维表格转 Markdown（首行作为表头）。

    - 清理单元格内换行（Markdown 表格要求单行）
    - 丢弃整行空行，补齐列数不齐的行（PDF 提取常见）
    - 无有效内容时返回空串
    """
    cleaned: List[List[str]] = []
    for row in rows or []:
        cells = [(c or "").replace("\n", " ").strip() for c in row]
        if any(cells):
            cleaned.append(cells)
    if not cleaned:
        return ""

    width = max(len(r) for r in cleaned)
    cleaned = [r + [""] * (width - len(r)) for r in cleaned]
    header, *body = cleaned

    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def _join_blocks(blocks: List[str]) -> str:
    """拼接文本块（正文与表格），块间以空行分隔"""
    return "\n\n".join(b for b in blocks if b.strip())


# ---------------------------------------------------------------------------
# 解析入口
# ---------------------------------------------------------------------------


def parse_file(file_path: str | Path) -> ParsedDocument:
    """
    按扩展名分派解析。

    Raises:
        ValueError: 不支持的格式
        FileNotFoundError: 文件不存在
    """
    path = Path(file_path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"不支持的文档格式: {ext}（支持 {sorted(SUPPORTED_EXTENSIONS)}）")
    if not path.exists():
        raise FileNotFoundError(f"文件不存在: {path}")

    if ext == ".pdf":
        return _parse_pdf(path)
    if ext == ".docx":
        return _parse_docx(path)
    file_type = "markdown" if ext in (".md", ".markdown") else "txt"
    return _parse_text_file(path, file_type)


def _parse_pdf(path: Path) -> ParsedDocument:
    """pdfplumber 逐页提取文本与表格，页间以空行分隔，保留每页文本供页码标注"""
    import pdfplumber  # 延迟导入：非 PDF 场景不加载该依赖

    pages: List[str] = []
    table_count = 0
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_text, tables = _extract_pdf_page(page)
            pages.append(page_text)
            table_count += tables
    text = "\n\n".join(pages)
    logger.info(
        "PDF 解析完成: %s，共 %d 页 / %d 表格 / %d 字符",
        path.name,
        len(pages),
        table_count,
        len(text),
    )
    return ParsedDocument(
        text=text,
        file_name=path.name,
        file_type="pdf",
        char_count=len(text),
        page_count=len(pages),
        pages=pages,
        table_count=table_count,
    )


def _extract_pdf_page(page) -> Tuple[str, int]:
    """
    提取单页文本 + 表格，返回 (合并后的文本, 表格数)。

    表格以 Markdown 追加在页面文本之后，使其参与切分与向量化。
    """
    blocks: List[str] = []
    text = page.extract_text() or ""
    if text:
        blocks.append(text)

    table_count = 0
    for rows in page.extract_tables() or []:
        md = table_to_markdown(rows)
        if md:
            table_count += 1
            blocks.append(f"{TABLE_MARKER.format(n=table_count)}\n{md}")
    return _join_blocks(blocks), table_count


def _parse_docx(path: Path) -> ParsedDocument:
    """python-docx 按文档顺序提取段落与表格，标题保留层级标记（# 数量）"""
    import docx  # 延迟导入

    document = docx.Document(str(path))
    blocks: List[str] = []
    table_count = 0
    para_count = 0
    for kind, payload in _iter_docx_blocks(document):
        if kind == "table":
            md = table_to_markdown(payload)  # type: ignore[arg-type]
            if md:
                table_count += 1
                blocks.append(f"{TABLE_MARKER.format(n=table_count)}\n{md}")
        else:
            text = str(payload).strip()
            if text:
                para_count += 1
                blocks.append(text)

    text = _join_blocks(blocks)
    logger.info(
        "Word 解析完成: %s，共 %d 段落 / %d 表格 / %d 字符",
        path.name,
        para_count,
        table_count,
        len(text),
    )
    return ParsedDocument(
        text=text,
        file_name=path.name,
        file_type="word",
        char_count=len(text),
        table_count=table_count,
    )


def _iter_docx_blocks(document) -> Iterator[Tuple[str, object]]:
    """
    按文档顺序产出 ("para", 文本) 与 ("table", 二维单元格) 。

    优先使用 python-docx 1.1+ 的 iter_inner_content（保持段落与表格的原始顺序），
    旧版本退化为先全部段落、再全部表格。
    """
    if hasattr(document, "iter_inner_content"):
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                yield "para", _docx_paragraph_text(block)
            elif isinstance(block, Table):
                yield "table", [[cell.text for cell in row.cells] for row in block.rows]
        return

    for para in document.paragraphs:
        yield "para", _docx_paragraph_text(para)
    for table in document.tables:
        yield "table", [[cell.text for cell in row.cells] for row in table.rows]


def _docx_paragraph_text(para) -> str:
    """段落文本，标题按层级加 Markdown 前缀（Heading 1 → #，Title → #）"""
    text = (para.text or "").strip()
    if not text:
        return ""
    style_name = (getattr(para.style, "name", "") or "")
    if style_name.startswith("Heading"):
        level = style_name.replace("Heading", "").strip()
        level = int(level) if level.isdigit() else 2
        return f"{'#' * min(level, 6)} {text}"
    if style_name == "Title":
        return f"# {text}"
    return text


def _parse_text_file(path: Path, file_type: str) -> ParsedDocument:
    """读取纯文本文件，UTF-8 优先，失败回退 GBK"""
    text = None
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            text = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError(f"无法识别文件编码（已尝试 utf-8/gbk）: {path.name}")
    return ParsedDocument(
        text=text,
        file_name=path.name,
        file_type=file_type,
        char_count=len(text),
    )
